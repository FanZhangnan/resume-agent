"""Word ZIP、正文事实、A4/字体配置与无效输入拒绝检查。"""

from _bootstrap import PROJECT_ROOT  # noqa: F401

from copy import deepcopy
from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from docx import Document

from mock_data import MOCK_SUGGESTIONS_ROUND2
from public_resume import normalize_public_resume
from resume_export import render_resume_docx


NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
W = "{" + NS["w"] + "}"


def _xml_parts(data):
    assert isinstance(data, bytes) and data.startswith(b"PK")
    with ZipFile(BytesIO(data)) as archive:
        assert archive.testzip() is None
        assert {"[Content_Types].xml", "word/document.xml", "word/styles.xml"}.issubset(archive.namelist())
        for name in archive.namelist():
            if name.endswith(".xml"):
                ET.fromstring(archive.read(name))
        return ET.fromstring(archive.read("word/document.xml")), ET.fromstring(archive.read("word/styles.xml"))


def _leaves(value):
    if isinstance(value, str):
        if value:
            yield value
    elif isinstance(value, dict):
        for child in value.values():
            yield from _leaves(child)
    elif isinstance(value, list):
        for child in value:
            yield from _leaves(child)


def test_docx_preserves_names_dates_and_every_normalized_fact():
    source = deepcopy(MOCK_SUGGESTIONS_ROUND2["optimized_resume_struct"])
    source["basic_info"]["name"] = "张南帆 Felix & <Fan>"
    source["basic_info"]["location"] = "Brisbane, QLD"
    source["education"][0]["school"] = "The University of Queensland 昆士兰大学"
    source["education"][0]["start"] = "2026-07"
    source["education"][0]["end"] = "2027-12（预计）"
    source["education"][0]["highlights"] = ["课程：DATA7703 & DATA7901"]
    source["projects"][0]["start"] = "2026.09"
    source["projects"][0]["end"] = "至今"
    source["extras"] = ["GitHub: https://github.com/FanZhangnan", "English: 未提供考试成绩"]
    source["private_credentials"] = "PRIVATE_SECRET"
    source["experience"][0]["private_note"] = "PRIVATE_SECRET"
    original = deepcopy(source)
    data = render_resume_docx(source)
    root, _ = _xml_parts(data)
    text = "\n".join(node.text or "" for node in root.findall(".//w:t", NS))
    for fact in _leaves(normalize_public_resume(source)):
        assert fact in text, fact
    assert "PRIVATE_SECRET" not in text
    assert "主导店铺全年30+场" not in text
    assert "一页" not in text
    reopened = Document(BytesIO(data))
    assert not reopened.tables
    assert source == original


def test_docx_has_a4_one_column_readable_fonts_and_real_list_style():
    root, styles = _xml_parts(render_resume_docx(MOCK_SUGGESTIONS_ROUND2["optimized_resume_struct"]))
    sections = root.findall(".//w:sectPr", NS)
    assert len(sections) == 1
    page = sections[0].find("w:pgSz", NS)
    assert abs(int(page.attrib[W + "w"]) - 11906) <= 1
    assert abs(int(page.attrib[W + "h"]) - 16838) <= 1
    assert sections[0].find("w:cols", NS).attrib[W + "num"] == "1"
    margins = sections[0].find("w:pgMar", NS)
    assert all(int(margins.attrib[W + side]) > 0 for side in ("top", "bottom", "left", "right"))
    assert not root.findall(".//w:tbl", NS)
    assert root.findall(".//w:pStyle[@w:val='ListBullet']", NS)
    for style_id in ("Normal", "Title", "Heading1", "ListBullet"):
        style = styles.find(f"w:style[@w:styleId='{style_id}']", NS)
        fonts = style.find("w:rPr/w:rFonts", NS)
        assert fonts.attrib[W + "eastAsia"] == "宋体"
        assert fonts.attrib[W + "ascii"] == "Arial"
        assert int(style.find("w:rPr/w:sz", NS).attrib[W + "val"]) >= 21


def test_minimal_resume_and_education_bullet_alias_remain_supported():
    source = {"education": [{"school": "学校", "end": "2026-12", "bullets": ["保留原始课程"]}]}
    root, _ = _xml_parts(render_resume_docx(source))
    text = "\n".join(node.text or "" for node in root.findall(".//w:t", NS))
    assert "学校" in text and "2026-12" in text and "保留原始课程" in text
    assert "候选人" not in text and "至今" not in text


def test_long_valid_content_is_not_truncated_to_fit_one_page():
    bullets = [str(index) + ":" + "有依据的详细经历。" * 100 for index in range(20)]
    source = {"experience": [{"company": "公司", "bullets": bullets}]}
    data = render_resume_docx(source)
    reopened = Document(BytesIO(data))
    exported = [paragraph.text for paragraph in reopened.paragraphs]
    assert all(bullet in exported for bullet in bullets)
    assert len([paragraph for paragraph in reopened.paragraphs if paragraph.style.name == "List Bullet"]) == 20


def test_invalid_resume_input_is_rejected():
    invalid_values = (None, [], "resume", {}, {"basic_info": {"name": "Only a name"}}, {"experience": "bad"}, {"experience": [False]}, {"education": [{"school": 123}]}, {"education": [{"school": "学校", "end": "中" * 501}]}, {"experience": [{"company": "公司", "bullets": [1]}]}, {"projects": [{"name": "P", "bullets": ["x" * 4001]}]}, {"experience": [{"company": "C"}] * 51})
    for source in invalid_values:
        try:
            render_resume_docx(source)
        except ValueError:
            continue
        raise AssertionError(f"Expected invalid input rejection: {source!r}")


if __name__ == "__main__":
    tests = [value for name, value in globals().copy().items() if name.startswith("test_")]
    for test in tests:
        test()
        print("PASS", test.__name__)
