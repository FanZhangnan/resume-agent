"""从规范化简历生成可编辑的 A4 Word；保留全部正文，不承诺页数。"""

from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from public_resume import normalize_public_resume


def _font(style, size, *, bold=False):
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")


def _record(document, record, title_fields, bullet_field):
    title = " | ".join(record[field] for field in title_fields if record[field])
    dates = " – ".join(record[field] for field in ("start", "end") if record[field])
    if title:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.keep_with_next = bool(dates or record[bullet_field])
        paragraph.add_run(title).bold = True
    if dates:
        paragraph = document.add_paragraph(dates)
        paragraph.paragraph_format.keep_with_next = bool(record[bullet_field])
    for text in record[bullet_field]:
        document.add_paragraph(text, style="List Bullet")


def render_resume_docx(resume_struct):
    """返回 DOCX 字节；无效简历抛出 ValueError，不写文件、不改写事实。"""
    resume = normalize_public_resume(resume_struct)
    if resume is None:
        raise ValueError("简历结构无效，无法导出 Word")

    document = Document()
    zoom = document.settings.element.find(qn("w:zoom"))
    if zoom is not None:
        zoom.set(qn("w:percent"), "100")
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = section.bottom_margin = Mm(16)
    section.left_margin = section.right_margin = Mm(18)
    section._sectPr.find(qn("w:cols")).set(qn("w:num"), "1")
    for name, size, bold in (
        ("Normal", 10.5, False), ("Title", 18, True),
        ("Heading 1", 12, True), ("List Bullet", 10.5, False),
    ):
        style = document.styles[name]
        _font(style, size, bold=bold)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.1
    document.styles["Heading 1"].paragraph_format.space_before = Pt(9)
    document.styles["Heading 1"].paragraph_format.keep_with_next = True

    basic = resume["basic_info"]
    if basic["name"]:
        document.add_paragraph(basic["name"], style="Title")
    contact = " | ".join(basic[field] for field in ("phone", "email", "location") if basic[field])
    if contact:
        document.add_paragraph(contact)
    if basic["target_role"]:
        document.add_paragraph("目标岗位：" + basic["target_role"])
    if resume["summary"]:
        document.add_heading("个人简介", level=1)
        document.add_paragraph(resume["summary"])
    for key, heading, fields, bullet_field in (
        ("education", "教育背景", ("school", "degree", "major"), "highlights"),
        ("experience", "工作经历", ("company", "title"), "bullets"),
        ("projects", "项目经历", ("name", "role"), "bullets"),
    ):
        if resume[key]:
            document.add_heading(heading, level=1)
            for record in resume[key]:
                _record(document, record, fields, bullet_field)
    if resume["skills"]:
        document.add_heading("技能", level=1)
        for skill in resume["skills"]:
            paragraph = document.add_paragraph()
            paragraph.add_run(skill["group"] + "：").bold = True
            paragraph.add_run("、".join(skill["items"]))
    if resume["extras"]:
        document.add_heading("其他", level=1)
        for text in resume["extras"]:
            document.add_paragraph(text, style="List Bullet")

    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()
